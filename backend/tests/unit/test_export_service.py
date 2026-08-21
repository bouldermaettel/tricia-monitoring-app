from io import BytesIO

from openpyxl import load_workbook

from src.api.schemas.cases import CaseCreateRequest
from src.services.case_service import CaseService
from src.services.export_service import DEFAULT_CASE_EXPORT_COLUMNS, ExportService


def test_export_csv_and_xlsx(db_session):
    CaseService(db_session).create_case(
        CaseCreateRequest(
            vk_number='VK-20260423-040',
            device_name='dev',
            tricia_s=1,
            tricia_p=1,
            tricia_d=5,
            user_s=1,
            user_d=5,
            validation_status='saved',
        ),
        actor_id='tester',
    )
    exporter = ExportService(db_session)
    assert b'vk_number' in exporter.to_csv()
    assert b'TRI-S' in exporter.to_csv()
    assert b'WIMI-S' in exporter.to_csv()
    workbook = load_workbook(BytesIO(exporter.to_xlsx()), read_only=True)
    assert list(next(workbook['cases'].iter_rows(values_only=True))) == [
        'id', 'vk_number', 'device_name', 'analysis_date', 'validation_status',
        'TRI-S', 'WIMI-S', 'TRI-P', 'WIMI-P', 'TRI-D', 'WIMI-D', 'TRI-RISK', 'WIMI-RISK',
    ]

    filtered_workbook = load_workbook(
        BytesIO(exporter.filtered_table_to_xlsx([], {})),
        read_only=True,
    )
    assert list(next(filtered_workbook['table'].iter_rows(values_only=True))) == DEFAULT_CASE_EXPORT_COLUMNS
